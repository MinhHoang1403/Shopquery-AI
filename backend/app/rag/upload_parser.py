"""Bounded parsing for untrusted dataset document uploads."""

from __future__ import annotations

import hashlib
import io
import zipfile
from dataclasses import dataclass
from pathlib import PurePath, PurePosixPath

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.errors import UploadValidationError
from app.rag.documents import MarkdownDocument, _clean_markdown

SUPPORTED_DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 500_000
MAX_PDF_PAGES = 100
MAX_DOCX_ENTRIES = 2_000
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 200


@dataclass(frozen=True)
class ParsedDocument:
    document: MarkdownDocument
    file_type: str
    content_hash: str


def parse_document_file(
    filename: str, content: bytes, *, max_bytes: int = MAX_DOCUMENT_BYTES
) -> ParsedDocument:
    safe_name = _safe_filename(filename)
    extension = PurePosixPath(safe_name).suffix.lower()
    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise UploadValidationError("Chỉ hỗ trợ file PDF, DOCX, TXT hoặc MD.")
    if not content:
        raise UploadValidationError("Tài liệu tải lên đang trống.")
    if len(content) > max_bytes:
        raise UploadValidationError(f"Tài liệu vượt quá giới hạn {max_bytes // (1024 * 1024)} MB.")

    if extension == ".pdf":
        title, extracted = _parse_pdf(content, PurePosixPath(safe_name).stem)
    elif extension == ".docx":
        title, extracted = _parse_docx(content, PurePosixPath(safe_name).stem)
    else:
        title, extracted = _parse_text(content, PurePosixPath(safe_name).stem)

    normalized = _clean_markdown(extracted)
    if not normalized:
        raise UploadValidationError("Tài liệu không có nội dung văn bản có thể đọc.")
    if not any(
        line.strip() and not line.lstrip().startswith("#")
        for line in normalized.splitlines()
    ):
        raise UploadValidationError("Tài liệu chỉ có tiêu đề, không có nội dung để hỏi đáp.")
    if len(normalized) > MAX_EXTRACTED_CHARACTERS:
        raise UploadValidationError("Nội dung trích xuất vượt quá giới hạn xử lý an toàn.")
    return ParsedDocument(
        document=MarkdownDocument(safe_name, title, normalized),
        file_type=extension.removeprefix("."),
        content_hash=hashlib.sha256(content).hexdigest(),
    )


def _parse_text(content: bytes, fallback_title: str) -> tuple[str, str]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise UploadValidationError("TXT/MD phải sử dụng encoding UTF-8.") from exc
    if "\x00" in text:
        raise UploadValidationError("Tài liệu chứa ký tự null không hợp lệ.")
    title = fallback_title
    for line in text.splitlines():
        if line.startswith("# ") and line[2:].strip():
            title = line[2:].strip()[:300]
            break
    return title, text


def _parse_pdf(content: bytes, fallback_title: str) -> tuple[str, str]:
    if not content.startswith(b"%PDF-"):
        raise UploadValidationError("Nội dung file không có chữ ký PDF hợp lệ.")
    try:
        reader = PdfReader(io.BytesIO(content), strict=True)
        if reader.is_encrypted:
            raise UploadValidationError("PDF mã hóa không được hỗ trợ.")
        if not reader.pages:
            raise UploadValidationError("PDF không có trang nào.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise UploadValidationError(f"PDF chỉ được có tối đa {MAX_PDF_PAGES} trang.")
        sections: list[str] = []
        extracted_size = 0
        for number, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if page_text:
                sections.append(f"## Trang {number}\n\n{page_text}")
                extracted_size += len(page_text)
            if extracted_size > MAX_EXTRACTED_CHARACTERS:
                raise UploadValidationError(
                    "Nội dung PDF trích xuất vượt quá giới hạn xử lý an toàn."
                )
    except UploadValidationError:
        raise
    except (PdfReadError, ValueError, TypeError, KeyError) as exc:
        raise UploadValidationError("PDF bị hỏng hoặc có cấu trúc không được hỗ trợ.") from exc
    if not sections:
        raise UploadValidationError(
            "PDF không có text layer; tài liệu cần OCR nhưng OCR chưa được hỗ trợ."
        )
    return fallback_title, f"# {fallback_title}\n\n" + "\n\n".join(sections)


def _parse_docx(content: bytes, fallback_title: str) -> tuple[str, str]:
    _validate_docx_archive(content)
    try:
        document = Document(io.BytesIO(content))
    except (PackageNotFoundError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise UploadValidationError("DOCX bị hỏng hoặc có cấu trúc không hợp lệ.") from exc

    lines: list[str] = []
    title = fallback_title
    extracted_size = 0
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if not value:
            continue
        style = (paragraph.style.name if paragraph.style else "").casefold()
        if style == "title":
            title = value[:300]
            lines.append(f"# {value}")
        elif style.startswith("heading"):
            level_text = style.removeprefix("heading").strip()
            level = int(level_text) if level_text.isdigit() else 2
            if level == 1 and title == fallback_title:
                title = value[:300]
            lines.append(f"{'#' * min(6, max(1, level))} {value}")
        else:
            lines.append(value)
        extracted_size += len(value)
        if extracted_size > MAX_EXTRACTED_CHARACTERS:
            raise UploadValidationError("Nội dung DOCX trích xuất vượt quá giới hạn xử lý an toàn.")
    return title, "\n\n".join(lines)


def _validate_docx_archive(content: bytes) -> None:
    if not content.startswith(b"PK"):
        raise UploadValidationError("Nội dung file không có chữ ký DOCX hợp lệ.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            names = {info.filename for info in infos}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise UploadValidationError("Archive không phải là tài liệu DOCX hợp lệ.")
            if len(infos) > MAX_DOCX_ENTRIES:
                raise UploadValidationError("DOCX có quá nhiều thành phần nén.")
            total_size = sum(info.file_size for info in infos)
            if total_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise UploadValidationError("DOCX giải nén vượt quá giới hạn an toàn.")
            for info in infos:
                path = PurePath(info.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise UploadValidationError("DOCX chứa đường dẫn archive không an toàn.")
                if info.flag_bits & 0x1:
                    raise UploadValidationError("DOCX mã hóa không được hỗ trợ.")
                if (
                    info.compress_size > 0
                    and info.file_size / info.compress_size > MAX_DOCX_COMPRESSION_RATIO
                ):
                    raise UploadValidationError("DOCX có tỷ lệ nén không an toàn.")
    except zipfile.BadZipFile as exc:
        raise UploadValidationError("DOCX không phải archive hợp lệ.") from exc


def _safe_filename(filename: str) -> str:
    normalized = (filename or "").replace("\\", "/")
    basename = PurePosixPath(normalized).name.strip()
    if not basename or basename in {".", ".."} or "\x00" in basename:
        raise UploadValidationError("Tên file không hợp lệ.")
    return basename[:255]
