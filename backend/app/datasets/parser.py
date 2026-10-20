"""Bounded CSV and XLSX parsing with conservative type inference."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath
from typing import Any
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from app.core.errors import UploadValidationError

SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 200
INTEGER_PATTERN = re.compile(r"^[+-]?\d+$")
DECIMAL_PATTERN = re.compile(r"^[+-]?\d+(?:\.\d+)?$")
ISO_DATE_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2}$")
ISO_DATETIME_PATTERN = re.compile(
    r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?(?:Z|[+-]\d{2}:?\d{2})?$"
)
TRUE_VALUES = {"true", "yes", "y", "1"}
FALSE_VALUES = {"false", "no", "n", "0"}


@dataclass(frozen=True)
class ColumnSpec:
    source_name: str
    name: str
    data_type: str

    def as_dict(self) -> dict[str, str]:
        return {
            "source_name": self.source_name,
            "name": self.name,
            "data_type": self.data_type,
        }


@dataclass(frozen=True)
class ParsedTable:
    filename: str
    file_type: str
    sheet_name: str | None
    sheet_names: tuple[str, ...]
    columns: tuple[ColumnSpec, ...]
    rows: tuple[tuple[Any, ...], ...]
    warnings: tuple[str, ...] = ()

    @property
    def row_count(self) -> int:
        return len(self.rows)

    @property
    def column_count(self) -> int:
        return len(self.columns)

    def preview(self, limit: int = 10) -> list[dict[str, Any]]:
        return [
            {
                column.source_name: _json_value(value)
                for column, value in zip(self.columns, row, strict=True)
            }
            for row in self.rows[:limit]
        ]

    def database_rows(self) -> list[dict[str, Any]]:
        return [
            {
                column.name: _convert_value(value, column.data_type)
                for column, value in zip(self.columns, row, strict=True)
            }
            for row in self.rows
        ]


def parse_tabular_file(
    filename: str,
    content: bytes,
    *,
    sheet_name: str | None = None,
    max_bytes: int,
    max_rows: int,
    max_columns: int,
) -> ParsedTable:
    safe_filename = _safe_filename(filename)
    extension = PurePosixPath(safe_filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise UploadValidationError("Chỉ hỗ trợ file CSV hoặc XLSX.")
    if not content:
        raise UploadValidationError("File tải lên đang trống.")
    if len(content) > max_bytes:
        raise UploadValidationError(
            f"File vượt quá giới hạn {max_bytes // (1024 * 1024)} MB."
        )
    if extension == ".csv":
        if sheet_name:
            raise UploadValidationError("CSV không hỗ trợ lựa chọn sheet.")
        return _parse_csv(safe_filename, content, max_rows, max_columns)
    return _parse_xlsx(safe_filename, content, sheet_name, max_rows, max_columns)


def _parse_csv(
    filename: str, content: bytes, max_rows: int, max_columns: int
) -> ParsedTable:
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise UploadValidationError("CSV phải sử dụng encoding UTF-8.") from exc
    if "\x00" in decoded:
        raise UploadValidationError("CSV chứa ký tự null không hợp lệ.")
    try:
        records = list(csv.reader(io.StringIO(decoded, newline=""), strict=True))
    except csv.Error as exc:
        raise UploadValidationError("CSV không đúng cấu trúc.") from exc
    if not records or not records[0]:
        raise UploadValidationError("CSV phải có hàng tiêu đề.")
    headers = _validate_headers(records[0], max_columns)
    rows: list[tuple[Any, ...]] = []
    formula_warning = False
    for line_number, record in enumerate(records[1:], start=2):
        if len(record) != len(headers):
            raise UploadValidationError(
                f"Dòng {line_number} có {len(record)} cột, cần đúng {len(headers)} cột."
            )
        if all(not value.strip() for value in record):
            continue
        if len(rows) >= max_rows:
            raise UploadValidationError(f"File vượt quá giới hạn {max_rows} dòng dữ liệu.")
        normalized = tuple(value if value != "" else None for value in record)
        formula_warning = formula_warning or any(
            isinstance(value, str) and value.startswith(("=", "+", "-", "@"))
            for value in normalized
        )
        rows.append(normalized)
    if not rows:
        raise UploadValidationError("File không có dòng dữ liệu nào.")
    columns = _build_columns(headers, rows)
    warnings = (
        ("Giá trị giống công thức được lưu dưới dạng dữ liệu, không được thực thi.",)
        if formula_warning
        else ()
    )
    return ParsedTable(filename, "csv", None, (), columns, tuple(rows), warnings)


def _parse_xlsx(
    filename: str,
    content: bytes,
    requested_sheet: str | None,
    max_rows: int,
    max_columns: int,
) -> ParsedTable:
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:
        raise UploadValidationError("XLSX không hợp lệ hoặc không thể đọc an toàn.") from exc
    try:
        sheet_names = tuple(workbook.sheetnames)
        if not sheet_names:
            raise UploadValidationError("XLSX không có sheet nào.")
        selected_sheet = requested_sheet or sheet_names[0]
        if selected_sheet not in sheet_names:
            raise UploadValidationError("Sheet đã chọn không tồn tại trong file.")
        worksheet = workbook[selected_sheet]
        dimension = worksheet.calculate_dimension(force=True)
        _, _, sheet_max_column, _ = range_boundaries(dimension)
        if sheet_max_column > max_columns:
            raise UploadValidationError(
                f"File vượt quá giới hạn {max_columns} cột dữ liệu."
            )
        row_iterator = worksheet.iter_rows(values_only=False)
        try:
            header_cells = next(row_iterator)
        except StopIteration as exc:
            raise UploadValidationError("XLSX phải có hàng tiêu đề.") from exc
        headers = _validate_headers([cell.value for cell in header_cells], max_columns)
        rows: list[tuple[Any, ...]] = []
        formula_warning = False
        for cells in row_iterator:
            values = tuple(_xlsx_cell_value(cell) for cell in cells[: len(headers)])
            if len(values) != len(headers):
                values += (None,) * (len(headers) - len(values))
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            if len(rows) >= max_rows:
                raise UploadValidationError(
                    f"File vượt quá giới hạn {max_rows} dòng dữ liệu."
                )
            for cell in cells[: len(headers)]:
                if cell.data_type == "f":
                    formula_warning = True
            rows.append(values)
        if not rows:
            raise UploadValidationError("File không có dòng dữ liệu nào.")
        columns = _build_columns(headers, rows)
        warnings = (
            ("Công thức Excel được lưu dưới dạng văn bản và không được thực thi.",)
            if formula_warning
            else ()
        )
        return ParsedTable(
            filename,
            "xlsx",
            selected_sheet,
            sheet_names,
            columns,
            tuple(rows),
            warnings,
        )
    except UploadValidationError:
        raise
    except ParseError as exc:
        raise UploadValidationError("XLSX không hợp lệ hoặc không thể đọc an toàn.") from exc
    finally:
        workbook.close()


def _xlsx_cell_value(cell: Any) -> Any:
    value = cell.value
    if isinstance(value, datetime):
        number_format = re.sub(r'"[^"]*"|\\.', "", cell.number_format.casefold())
        if "h" not in number_format and "s" not in number_format:
            return value.date()
    return value


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if not infos:
                raise UploadValidationError("XLSX không có nội dung.")
            total_size = sum(info.file_size for info in infos)
            if total_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise UploadValidationError("XLSX giải nén vượt quá giới hạn an toàn.")
            for info in infos:
                if info.flag_bits & 0x1:
                    raise UploadValidationError("XLSX mã hóa không được hỗ trợ.")
                if (
                    info.compress_size > 0
                    and info.file_size / info.compress_size > MAX_XLSX_COMPRESSION_RATIO
                ):
                    raise UploadValidationError("XLSX có tỷ lệ nén không an toàn.")
    except zipfile.BadZipFile as exc:
        raise UploadValidationError("XLSX không phải archive hợp lệ.") from exc


def _safe_filename(filename: str) -> str:
    normalized = (filename or "").replace("\\", "/")
    basename = PurePosixPath(normalized).name.strip()
    if not basename or basename in {".", ".."} or "\x00" in basename:
        raise UploadValidationError("Tên file không hợp lệ.")
    return basename[:255]


def _validate_headers(raw_headers: list[Any], max_columns: int) -> tuple[str, ...]:
    if not raw_headers or len(raw_headers) > max_columns:
        raise UploadValidationError(f"File chỉ được có tối đa {max_columns} cột.")
    headers = tuple("" if value is None else str(value).strip() for value in raw_headers)
    if any(not header for header in headers):
        raise UploadValidationError("Tên cột không được để trống.")
    folded = [header.casefold() for header in headers]
    if len(folded) != len(set(folded)):
        raise UploadValidationError("Tên cột không được trùng nhau.")
    return headers


def _build_columns(
    headers: tuple[str, ...], rows: list[tuple[Any, ...]]
) -> tuple[ColumnSpec, ...]:
    used_names: set[str] = set()
    columns: list[ColumnSpec] = []
    for index, header in enumerate(headers):
        base_name = _identifier(header, index + 1)
        name = base_name
        suffix = 2
        while name in used_names:
            name = f"{base_name}_{suffix}"
            suffix += 1
        used_names.add(name)
        values = [row[index] for row in rows if row[index] not in (None, "")]
        try:
            data_type = _infer_type(values)
        except (ValueError, UploadValidationError) as exc:
            raise UploadValidationError(f'Cột "{header}": {exc}') from exc
        columns.append(ColumnSpec(header, name, data_type))
    return tuple(columns)


def _identifier(header: str, position: int) -> str:
    header = header.translate(str.maketrans({"Đ": "D", "đ": "d"}))
    ascii_name = (
        unicodedata.normalize("NFKD", header).encode("ascii", "ignore").decode("ascii")
    )
    identifier = re.sub(r"[^a-zA-Z0-9]+", "_", ascii_name).strip("_").lower()
    if not identifier or identifier[0].isdigit():
        identifier = f"column_{position}"
    return identifier[:50]


def _infer_type(values: list[Any]) -> str:
    if not values:
        return "TEXT"
    if any(_has_leading_zero(value) for value in values):
        return "TEXT"
    if all(isinstance(value, bool) or _is_boolean_text(value) for value in values):
        return "BOOLEAN"
    if all(_is_integer(value) for value in values):
        if any(not -(2**63) <= int(value) < 2**63 for value in values):
            raise UploadValidationError("Số nguyên vượt giới hạn BIGINT 64-bit.")
        return "BIGINT"
    if all(_is_decimal(value) for value in values):
        return "NUMERIC"
    datetimes = [_parse_datetime(value) for value in values if _is_datetime(value)]
    dates = [value for value in values if _is_date(value)]
    awareness = {value.utcoffset() is not None for value in datetimes}
    if len(awareness) > 1:
        raise UploadValidationError(
            "Không được trộn datetime có offset và không có offset. Hãy thống nhất kiểu thời gian."
        )
    if datetimes and dates:
        raise UploadValidationError(
            "Không được trộn DATE và datetime. Hãy thống nhất kiểu thời gian."
        )
    if len(datetimes) == len(values):
        return "TIMESTAMPTZ" if True in awareness else "TIMESTAMP"
    if len(dates) == len(values):
        for value in dates:
            if isinstance(value, str):
                date.fromisoformat(value.strip())
        return "DATE"
    return "TEXT"


def _has_leading_zero(value: Any) -> bool:
    if not isinstance(value, str):
        return False
    stripped = value.strip().lstrip("+-")
    return len(stripped) > 1 and stripped.startswith("0") and stripped.isdigit()


def _is_boolean_text(value: Any) -> bool:
    return isinstance(value, str) and value.strip().casefold() in TRUE_VALUES | FALSE_VALUES


def _is_integer(value: Any) -> bool:
    return isinstance(value, int) and not isinstance(value, bool) or (
        isinstance(value, str) and bool(INTEGER_PATTERN.fullmatch(value.strip()))
    )


def _is_decimal(value: Any) -> bool:
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return True
    return isinstance(value, str) and bool(DECIMAL_PATTERN.fullmatch(value.strip()))


def _is_datetime(value: Any) -> bool:
    return isinstance(value, datetime) or (
        isinstance(value, str) and bool(ISO_DATETIME_PATTERN.fullmatch(value.strip()))
    )


def _is_date(value: Any) -> bool:
    return isinstance(value, date) and not isinstance(value, datetime) or (
        isinstance(value, str) and bool(ISO_DATE_PATTERN.fullmatch(value.strip()))
    )


def _parse_datetime(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value
    return datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))


def _convert_value(value: Any, data_type: str) -> Any:
    if value is None or value == "":
        return None
    if data_type == "TEXT":
        return _json_value(value)
    if data_type == "BIGINT":
        return int(value)
    if data_type == "NUMERIC":
        try:
            return Decimal(str(value))
        except InvalidOperation as exc:
            raise UploadValidationError("Không thể chuyển giá trị NUMERIC.") from exc
    if data_type == "BOOLEAN":
        if isinstance(value, bool):
            return value
        return str(value).strip().casefold() in TRUE_VALUES
    if data_type == "DATE":
        return value if isinstance(value, date) else date.fromisoformat(str(value).strip())
    if data_type in {"TIMESTAMP", "TIMESTAMPTZ"}:
        converted = _parse_datetime(value)
        if (converted.utcoffset() is not None) != (data_type == "TIMESTAMPTZ"):
            raise UploadValidationError("Datetime không khớp kiểu timezone đã xác nhận.")
        return converted
    raise UploadValidationError(f"Kiểu dữ liệu không được hỗ trợ: {data_type}")


def _json_value(value: Any) -> Any:
    if isinstance(value, int) and abs(value) > 9_007_199_254_740_991:
        return str(value)
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value
