"""Execution-based golden and metamorphic coverage, never external model evaluation."""

import json
import re
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document
from sqlalchemy import text

from app.core.config import get_settings
from app.core.diagnostics import diagnostic_request
from app.database.session import get_readonly_engine, get_setup_engine
from app.datasets.relationship_repository import RelationshipRepository
from app.datasets.relationships import RelationshipConfirm, RelationshipInput
from app.datasets.repository import DatasetRepository
from app.dependencies import build_dataset_chat_service, build_workspace_context
from app.rag.pipeline import RagAnswerDraft
from app.rag.upload_parser import parse_document_file
from app.routing.hybrid import HybridAnswerDraft
from app.routing.router import RouteDecision, RouteLabel
from app.text_to_sql.query_plan import PlanDecision
from tests.auth_helpers import TEST_OWNER
from tests.hybrid_fixtures import (
    assert_decision_summary,
    classification_draft,
    prose_policy,
    rule_block,
)
from tests.join_domains import DOMAINS, uploads
from tests.plan_fixtures import decision_from_sql
from tests.test_dataset_parser import parse
from tests.test_guest_integration import OfflineEmbedder

pytestmark = pytest.mark.integration


@pytest.fixture
def domain_workspace(request):
    domain, seed = request.param
    repository = DatasetRepository(get_setup_engine(), get_readonly_engine(), TEST_OWNER)
    dataset = repository.create_workspace("Golden synthetic " + uuid4().hex)
    relations = RelationshipRepository(get_setup_engine(), TEST_OWNER)
    files, facts = uploads(domain, seed)
    try:
        for filename, content in files:
            dataset = repository.add_table(
                dataset.dataset_id, parse(filename, content), auto_relationships=True
            )
        tables = {t.display_name: t for t in dataset.tables}
        outer, inner, fact = (tables[n] for n in domain.labels)
        requests = [
            RelationshipInput(
                source_table_id=str(fact.dataset_table_id),
                source_column=domain.keys[1],
                target_table_id=str(inner.dataset_table_id),
                target_column=domain.keys[1],
            ),
            RelationshipInput(
                source_table_id=str(
                    fact.dataset_table_id if domain.star else inner.dataset_table_id
                ),
                source_column=domain.keys[0],
                target_table_id=str(outer.dataset_table_id),
                target_column=domain.keys[0],
            ),
        ]
        existing = {
            (r.source_table_id, r.target_table_id) for r in relations.list(dataset.dataset_id)
        }
        for candidate in requests:
            if (candidate.source_table_id, candidate.target_table_id) not in existing:
                preview = relations.preview(dataset.dataset_id, candidate)
                assert preview.valid
                relations.confirm(
                    dataset.dataset_id,
                    RelationshipConfirm(
                        **candidate.model_dump(), revision=preview.revision, confirmed=True
                    ),
                )
        settings = get_settings().model_copy(update={"sql_join_v1_enabled": True})
        context = build_workspace_context(settings, dataset, repository)
        assert len(context.approved_joins) == 2
        yield domain, seed, repository, dataset, settings, context, facts
    finally:
        repository.delete(dataset.dataset_id, delete_conversations=True)


def golden_sql(domain, dataset, tables_count, filtered=False):
    names = {t.display_name: t.physical_table_name for t in dataset.tables}
    outer, inner, fact = (names[n] for n in domain.labels)
    k, d = domain.keys
    bindings = f"FROM uploaded_data.{fact} f"
    identity, label = f"f.{d}", None
    if tables_count >= 2:
        bindings += f" INNER JOIN uploaded_data.{inner} d ON f.{d}=d.{d}"
        identity, label = f"d.{d}", f"d.{domain.names[1]}"
    if tables_count == 3:
        bindings += (
            f" INNER JOIN uploaded_data.{outer} k ON {'f' if domain.star else 'd'}.{k}=k.{k}"
        )
        identity, label = f"k.{k}", f"k.{domain.names[0]}"
    selection = identity + " AS entity_id" + (f",{label} AS entity_name" if label else "")
    group = identity + (f",{label}" if label else "")
    where = (
        " WHERE f.recorded_at >= TIMESTAMPTZ '2025-09-01 00:00:00+07:00' "
        "AND f.recorded_at < TIMESTAMPTZ '2025-10-01 00:00:00+07:00'"
        if filtered
        else ""
    )
    return (
        f"SELECT {selection},SUM(f.{domain.measure}) AS total,AVG(f.{domain.measure}) AS average,"
        f"COUNT(*) AS fact_row_count,COUNT(f.{domain.measure}) AS fact_value_count "
        + bindings
        + where
        + f" GROUP BY {group} ORDER BY entity_id"
    )


def oracle(facts, tables_count, filtered=False):
    parameters, values = {}, []
    for index, row in enumerate(facts):
        if filtered and row[5].startswith("2025-09-30"):
            continue
        identity, label = (row[2], row[3]) if tables_count == 3 else (row[0], row[1])
        parameters.update({f"id{index}": identity, f"name{index}": label, f"v{index}": row[4]})
        values.append(f"(:id{index},:name{index},CAST(:v{index} AS NUMERIC))")
    name = ",entity_name" if tables_count >= 2 else ""
    query = (
        f"SELECT entity_id{name},SUM(v)::text AS total,AVG(v)::text AS average,"
        "COUNT(*) AS fact_row_count,COUNT(v) AS fact_value_count "
        f"FROM (VALUES {','.join(values)}) oracle(entity_id,entity_name,v) "
        f"GROUP BY entity_id{name} ORDER BY entity_id"
    )
    with get_setup_engine().connect() as connection:
        return [dict(r) for r in connection.execute(text(query), parameters).mappings()]


PARAMETERS = [(domain, seed) for domain in DOMAINS for seed in (0, 7, 19)]


@pytest.mark.parametrize(
    "domain_workspace",
    PARAMETERS,
    indirect=True,
    ids=[f"{d.name}-{seed}" for d, seed in PARAMETERS],
)
@pytest.mark.parametrize("model", ["gemini", "inetcloud-knowledge"])
@pytest.mark.parametrize("tables_count", [1, 2, 3])
def test_domain_golden_results_and_metamorphic_uploads(
    domain_workspace, monkeypatch, model, tables_count
):
    domain, seed, _, dataset, settings, context, facts = domain_workspace
    sql = golden_sql(domain, dataset, tables_count, filtered=domain.name == "energy")
    expected = oracle(facts, tables_count, filtered=domain.name == "energy")
    calls = []

    class Gateway:
        def generate_structured(self, prompt, schema):
            calls.append(schema.__name__)
            assert all(t.physical_table_name in prompt for t in dataset.tables)
            if schema is RouteDecision:
                return RouteDecision(route=RouteLabel.SQL, reason="Golden offline")
            assert schema is PlanDecision
            return decision_from_sql(sql, prompt)

        def generate_text(self, prompt):
            calls.append("answer")
            rows = json.loads(prompt.split("Query result: ", 1)[1].split("\nJSON values", 1)[0])
            assert rows == expected
            return json.dumps(rows, ensure_ascii=False)

    monkeypatch.setattr("app.ai.chat.build_generation_gateway", lambda *_: Gateway())
    with diagnostic_request() as trace:
        result = build_dataset_chat_service(
            settings.model_copy(update={"ai_default_model": model}), context
        ).chat(domain.question)
    assert (
        result.route == RouteLabel.SQL and result.used_model == model and not result.fallback_used
    )
    assert result.rows == expected and result.answer
    assert result.sql.count("JOIN") == tables_count - 1
    assert calls == ["RouteDecision", "PlanDecision"]
    assert trace.events[-1]["sql_validated"] and trace.events[-1]["execution_reached"]
    assert trace.events[-1]["result_rows"] == len(expected)
    if domain.name == "logistics":
        assert sum(r["fact_row_count"] for r in result.rows) == 50
        assert sum(Decimal(r["total"]) for r in result.rows) == Decimal("2472500.00")
        if tables_count == 3:
            assert len(result.rows) == 5 and all(r["fact_row_count"] == 10 for r in result.rows)


@pytest.mark.parametrize(
    "domain_workspace", [(d, 7) for d in DOMAINS], indirect=True, ids=[d.name for d in DOMAINS]
)
@pytest.mark.parametrize("model", ["gemini", "inetcloud-knowledge"])
@pytest.mark.parametrize("route", [RouteLabel.RAG, RouteLabel.HYBRID, RouteLabel.UNSUPPORTED])
def test_domain_document_routing_and_grounding(domain_workspace, monkeypatch, model, route):
    from app.dependencies import build_document_repository

    domain, _, repository, dataset, settings, _, facts = domain_workspace
    embedder = OfflineEmbedder()
    documents = build_document_repository(TEST_OWNER)
    filename = "synthetic_kpi.md"
    prose = prose_policy(domain.keys[0], domain.measure, "40.00")
    content = ("# Synthetic KPI\n\n" + prose).encode()
    content += ("\n" + rule_block(domain.keys[0], domain.measure, "40.00")).encode()
    if domain.name == "logistics":
        filename = "synthetic_kpi.docx"
        policy = Document()
        policy.add_heading("Synthetic KPI", level=1)
        policy.add_table(rows=1, cols=1).cell(0, 0).text = prose
        policy.add_paragraph(rule_block(domain.keys[0], domain.measure, "40.00"))
        output = BytesIO()
        policy.save(output)
        content = output.getvalue()
    document = parse_document_file(filename, content, max_bytes=100_000)
    documents.add_document(dataset.dataset_id, document, embedder)
    dataset = repository.get(dataset.dataset_id)
    context = build_workspace_context(settings, dataset, repository)
    expected = oracle(facts, 3)
    qualified = [
        r["entity_id"]
        for r in expected
        if r["total"] is not None and Decimal(r["total"]) >= Decimal("40.00")
    ]
    statuses = [
        "insufficient_evidence"
        if r["total"] is None
        else "qualified"
        if r["entity_id"] in qualified
        else "not_qualified"
        for r in expected
    ]
    calls = []

    class Gateway:
        def generate_structured(self, prompt, schema):
            calls.append(schema.__name__)
            if schema is RouteDecision:
                return RouteDecision(route=route, reason="Golden document evidence")
            if schema is PlanDecision:
                return decision_from_sql(golden_sql(domain, dataset, 3), prompt)
            ids = [int(value) for value in re.findall(r"\[chunk_id=(\d+)", prompt)]
            assert ids and "40.00" in prompt
            if schema is RagAnswerDraft:
                return RagAnswerDraft(
                    answer="Synthetic threshold 40.00, all statuses", used_chunk_ids=[ids[0]]
                )
            assert schema is HybridAnswerDraft
            rows = json.loads(prompt.split("SQL rows: ", 1)[1].split("\n", 1)[0])
            assert rows == expected
            return classification_draft(expected, statuses, ids[0])

        def generate_text(self, prompt):
            return json.dumps(expected)

    monkeypatch.setattr("app.ai.chat.build_generation_gateway", lambda *_: Gateway())
    monkeypatch.setattr("app.dependencies.build_embedding_gateway", lambda _: embedder)
    result = build_dataset_chat_service(
        settings.model_copy(update={"ai_default_model": model}), context
    ).chat("Synthetic question with KPI evidence")
    assert result.route == route and result.used_model == model and not result.fallback_used
    if route == RouteLabel.HYBRID:
        assert result.rows == expected
        assert_decision_summary(result.answer, expected, statuses)
    else:
        assert result.rows == [] and "PlanDecision" not in calls
    if route != RouteLabel.UNSUPPORTED:
        assert len(result.sources) == 1 and result.sources[0].source_file == filename
    else:
        assert result.sources == [] and calls == ["RouteDecision"]
