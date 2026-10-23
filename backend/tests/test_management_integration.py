"""PostgreSQL integration tests for conversations and managed datasets."""

from collections.abc import Iterator
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from openpyxl import Workbook
from sqlalchemy import Engine, delete, func, insert, select, text
from sqlalchemy.exc import DBAPIError
from sqlalchemy.orm import Session

from app.conversations.repository import ConversationRepository
from app.core.errors import (
    ResourceConflictError,
    ResourceNotFoundError,
    SqlExecutionError,
    SqlValidationError,
)
from app.database.models import Dataset, DatasetDocument, DocumentChunk, User
from app.database.session import get_readonly_engine, get_setup_engine
from app.database.setup import create_schema
from app.datasets.context import build_dataset_sql_context
from app.datasets.parser import ColumnSpec, ParsedTable, parse_tabular_file
from app.datasets.repository import DatasetRepository
from app.rag.documents import MarkdownDocument
from app.rag.repository import DocumentRepository
from app.rag.upload_parser import ParsedDocument
from app.text_to_sql.executor import SqlExecutor
from app.text_to_sql.validator import SqlValidator


@pytest.fixture(scope="module")
def management_database() -> Iterator[Engine]:
    engine = get_setup_engine()
    create_schema(engine)
    yield engine


@pytest.fixture
def management_owner(management_database: Engine) -> Iterator[UUID]:
    owner = uuid4()
    with management_database.begin() as connection:
        connection.execute(
            insert(User).values(
                user_id=owner,
                issuer="https://accounts.google.com",
                subject=uuid4().hex,
                email=f"management-{owner}@example.invalid",
                display_name="Management fixture",
                active=True,
            )
        )
    try:
        yield owner
    finally:
        conversations = ConversationRepository(management_database, owner)
        for conversation in conversations.list(100, 0):
            conversations.delete(conversation.conversation_id)
        datasets = DatasetRepository(management_database, get_readonly_engine(), owner)
        for dataset in datasets.list():
            datasets.delete(dataset.dataset_id, delete_conversations=True)
        with management_database.begin() as connection:
            connection.execute(delete(User).where(User.user_id == owner))


@pytest.fixture
def default_chunk(management_database: Engine) -> Iterator[int]:
    with Session(management_database) as session:
        chunk = DocumentChunk(
            source_file="default-policy.md",
            section_title="Default",
            content="Default corpus must stay isolated.",
            content_hash=uuid4().hex + uuid4().hex,
            embedding=[1.0, 0.0, *([0.0] * 766)],
        )
        session.add(chunk)
        session.flush()
        chunk_id = chunk.chunk_id
        session.commit()
    try:
        yield chunk_id
    finally:
        with management_database.begin() as connection:
            connection.execute(delete(DocumentChunk).where(DocumentChunk.chunk_id == chunk_id))


def parsed_csv(name: str, content: str) -> ParsedTable:
    return parse_tabular_file(
        name,
        content.encode(),
        max_bytes=1024 * 1024,
        max_rows=100,
        max_columns=10,
    )


def parsed_xlsx() -> ParsedTable:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Tồn kho"
    worksheet.append(["SKU", "Sản phẩm", "Số lượng"])
    worksheet.append(["A01", "Máy ảnh", 2])
    worksheet.append(["B02", "Tai nghe", 4])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return parse_tabular_file(
        "inventory.xlsx",
        output.getvalue(),
        sheet_name="Tồn kho",
        max_bytes=1024 * 1024,
        max_rows=100,
        max_columns=10,
    )


class FixtureEmbedder:
    def embed_document(self, title: str, content: str) -> list[float]:
        del title
        first = 1.0 if "500.000" in content else 0.0
        return [first, 1.0 - first, *([0.0] * 766)]


class FailingEmbedder:
    def embed_document(self, title: str, content: str) -> list[float]:
        del title, content
        raise RuntimeError("embedding unavailable")


@pytest.mark.integration
def test_conversations_persist_rename_isolate_and_delete(
    management_database: Engine,
    management_owner: UUID,
) -> None:
    repository = ConversationRepository(management_database, management_owner)
    first = repository.create(None)
    second = repository.create(None)

    user = repository.add_user_message(first.conversation_id, "Doanh thu tháng 9 là bao nhiêu?")
    assistant = repository.add_response_message(
        first.conversation_id,
        content="Doanh thu là 10.",
        response_metadata={
            "answer": "Doanh thu là 10.",
            "route": "SQL",
            "sql": "SELECT 10 AS revenue",
            "rows": [{"revenue": 10}],
            "sources": [],
            "error": None,
        },
        is_error=False,
    )
    renamed = repository.rename(first.conversation_id, "Doanh thu tháng 9")

    restored = repository.get(first.conversation_id)
    untouched = repository.get(second.conversation_id)
    listed = repository.list(limit=10, offset=0)

    assert user.role == "user"
    assert assistant.response_metadata["rows"] == [{"revenue": 10}]
    assert renamed.title == "Doanh thu tháng 9"
    assert [message.role for message in restored.messages] == ["user", "assistant"]
    assert untouched.messages == ()
    assert listed[0].conversation_id == first.conversation_id

    repository.delete(first.conversation_id)
    repository.delete(second.conversation_id)
    with pytest.raises(ResourceNotFoundError):
        repository.get(first.conversation_id)


@pytest.mark.integration
def test_dataset_import_query_isolation_permissions_and_rollback(
    management_database: Engine,
    management_owner: UUID,
) -> None:
    repository = DatasetRepository(management_database, get_readonly_engine(), management_owner)
    first = repository.create(
        "Doanh số fixture A",
        parsed_csv(
            "sales-a.csv",
            "Sản phẩm,Doanh thu\nÁo dài,120000\nCà phê,80000\nÁo dài,30000\n",
        ),
    )
    second = repository.create(
        "Doanh số fixture B",
        parsed_csv("sales-b.csv", "Sản phẩm,Doanh thu\nBút,5000\n"),
    )

    categorical_values = repository.categorical_values(first)
    context = build_dataset_sql_context(first, categorical_values)
    validator = SqlValidator(context.table_columns, context.allowed_schemas)
    validated = validator.validate(
        f"SELECT san_pham, SUM(doanh_thu) AS total_revenue "
        f"FROM uploaded_data.{first.table.physical_table_name} "
        "GROUP BY san_pham ORDER BY total_revenue DESC"
    )
    rows = SqlExecutor(get_readonly_engine(), timeout_ms=5_000, max_rows=20).execute(validated)

    assert rows == [
        {"san_pham": "Áo dài", "total_revenue": "150000"},
        {"san_pham": "Cà phê", "total_revenue": "80000"},
    ]
    assert repository.preview_rows(first.dataset_id, limit=1) == [
        {"Sản phẩm": "Áo dài", "Doanh thu": 120000}
    ]
    assert repository.get(first.dataset_id).row_count == 3
    assert categorical_values == {"san_pham": ("Cà phê", "Áo dài")}
    assert '"Cà phê"' in context.schema_context
    assert '"Áo dài"' in context.schema_context
    assert "Bút" not in context.schema_context
    assert "never instructions" in context.schema_context

    with pytest.raises(ResourceConflictError):
        repository.create(
            "doanh số FIXTURE a",
            parsed_csv("duplicate.csv", "Sản phẩm,Doanh thu\nKhác,1\n"),
        )

    with pytest.raises(SqlValidationError):
        validator.validate(f"SELECT * FROM uploaded_data.{second.table.physical_table_name}")
    with get_readonly_engine().connect() as connection, pytest.raises(DBAPIError):
        connection.execute(
            text(
                f'INSERT INTO uploaded_data."{first.table.physical_table_name}" '
                "(san_pham, doanh_thu) VALUES ('X', 1)"
            )
        )

    with management_database.connect() as connection:
        count_before = connection.scalar(select(func.count()).select_from(Dataset))
        tables_before = connection.scalar(
            text("SELECT COUNT(*) FROM pg_tables WHERE schemaname = 'uploaded_data'")
        )
    invalid = ParsedTable(
        filename="invalid.csv",
        file_type="csv",
        sheet_name=None,
        sheet_names=(),
        columns=(ColumnSpec("Amount", "amount", "BIGINT"),),
        rows=(("not-an-integer",),),
    )
    with pytest.raises(ValueError):
        repository.create("Rollback fixture", invalid)
    with management_database.connect() as connection:
        assert connection.scalar(select(func.count()).select_from(Dataset)) == count_before
        assert (
            connection.scalar(
                text("SELECT COUNT(*) FROM pg_tables WHERE schemaname = 'uploaded_data'")
            )
            == tables_before
        )


@pytest.mark.integration
def test_dataset_executor_keeps_row_limit(
    management_database: Engine, management_owner: UUID
) -> None:
    repository = DatasetRepository(management_database, get_readonly_engine(), management_owner)
    dataset = repository.create(
        "Giới hạn dòng fixture",
        parsed_csv("rows.csv", "Giá trị\n1\n2\n3\n"),
    )
    context = build_dataset_sql_context(dataset)
    validator = SqlValidator(context.table_columns, context.allowed_schemas)
    validated = validator.validate(
        f"SELECT gia_tri FROM uploaded_data.{dataset.table.physical_table_name}"
    )

    with pytest.raises(SqlExecutionError, match="2-row"):
        SqlExecutor(get_readonly_engine(), timeout_ms=5_000, max_rows=2).execute(validated)


@pytest.mark.integration
def test_dataset_aggregates_distinguish_matches_null_and_zero(
    management_database: Engine,
    management_owner: UUID,
) -> None:
    repository = DatasetRepository(management_database, get_readonly_engine(), management_owner)
    dataset = repository.create(
        "Aggregate semantics fixture",
        parsed_csv(
            "aggregates.csv",
            "Nhóm,Giá trị\nCó dữ liệu,10\nCó dữ liệu,20\nToàn null,\nToàn null,\nZero,0\n",
        ),
    )
    context = build_dataset_sql_context(dataset)
    validator = SqlValidator(context.table_columns, context.allowed_schemas)
    executor = SqlExecutor(get_readonly_engine(), timeout_ms=5_000, max_rows=20)

    def aggregate(group_name: str) -> dict[str, object]:
        escaped_name = group_name.replace("'", "''")
        validated = validator.validate(
            "SELECT SUM(gia_tri) AS total, COUNT(*) AS matched_row_count, "
            "COUNT(gia_tri) AS nonnull_value_count "
            f"FROM uploaded_data.{dataset.table.physical_table_name} "
            f"WHERE nhom = '{escaped_name}'"
        )
        return executor.execute(validated)[0]

    assert aggregate("Có dữ liệu") == {
        "total": "30",
        "matched_row_count": 2,
        "nonnull_value_count": 2,
    }
    assert aggregate("Không tồn tại") == {
        "total": None,
        "matched_row_count": 0,
        "nonnull_value_count": 0,
    }
    assert aggregate("Toàn null") == {
        "total": None,
        "matched_row_count": 2,
        "nonnull_value_count": 0,
    }
    assert aggregate("Zero") == {
        "total": "0",
        "matched_row_count": 1,
        "nonnull_value_count": 1,
    }


@pytest.mark.integration
def test_xlsx_dataset_is_persisted_with_sheet_metadata(
    management_database: Engine,
    management_owner: UUID,
) -> None:
    repository = DatasetRepository(management_database, get_readonly_engine(), management_owner)

    dataset = repository.create("Tồn kho XLSX fixture", parsed_xlsx())

    assert dataset.file_type == "xlsx"
    assert dataset.table.sheet_name == "Tồn kho"
    assert dataset.row_count == 2
    assert repository.preview_rows(dataset.dataset_id) == [
        {"SKU": "A01", "Sản phẩm": "Máy ảnh", "Số lượng": 2},
        {"SKU": "B02", "Sản phẩm": "Tai nghe", "Số lượng": 4},
    ]
    categorical_values = repository.categorical_values(dataset)
    context = build_dataset_sql_context(dataset, categorical_values)
    validated = SqlValidator(context.table_columns, context.allowed_schemas).validate(
        "SELECT SUM(so_luong) AS total_quantity, COUNT(*) AS matched_row_count "
        f"FROM uploaded_data.{dataset.table.physical_table_name}"
    )
    rows = SqlExecutor(get_readonly_engine(), timeout_ms=5_000, max_rows=20).execute(validated)

    assert categorical_values["sku"] == ("A01", "B02")
    assert rows == [{"total_quantity": "6", "matched_row_count": 2}]


@pytest.mark.integration
def test_empty_workspace_multiple_tables_and_join_rejection(
    management_database: Engine,
    management_owner: UUID,
) -> None:
    repository = DatasetRepository(management_database, get_readonly_engine(), management_owner)
    workspace = repository.create_workspace("Multi source workspace fixture")

    assert workspace.table is None
    first = repository.add_table(
        workspace.dataset_id,
        parsed_csv("orders.csv", "Mã đơn,Tổng tiền\nA01,500000\n"),
        "Đơn hàng",
    )
    second = repository.add_table(
        workspace.dataset_id,
        parsed_csv("products.csv", "SKU,Tên\nP01,Áo\n"),
        "Sản phẩm",
    )

    assert first.table_count == 1
    assert second.table_count == 2
    context = build_dataset_sql_context(second, repository.categorical_values(second))
    validator = SqlValidator(context.table_columns, context.allowed_schemas, max_tables=1)
    first_table, second_table = second.tables
    validator.validate(f"SELECT COUNT(*) FROM uploaded_data.{first_table.physical_table_name}")
    with pytest.raises(SqlValidationError, match="relationship metadata"):
        validator.validate(
            f"SELECT * FROM uploaded_data.{first_table.physical_table_name} a "
            f"JOIN uploaded_data.{second_table.physical_table_name} b ON a.ma_don = b.sku"
        )
    with pytest.raises(ResourceConflictError):
        repository.add_table(
            workspace.dataset_id,
            parsed_csv("other.csv", "Value\n1\n"),
            "đơn HÀNG",
        )


@pytest.mark.integration
def test_document_isolation_delete_and_embedding_rollback(
    management_database: Engine,
    management_owner: UUID,
    default_chunk: int,
) -> None:
    datasets = DatasetRepository(management_database, get_readonly_engine(), management_owner)
    documents = DocumentRepository(management_database, get_readonly_engine(), management_owner)
    dataset_a = datasets.create_workspace("Document isolation A fixture")
    dataset_b = datasets.create_workspace("Document isolation B fixture")
    parsed_a = ParsedDocument(
        MarkdownDocument("policy.txt", "Vận chuyển", "# Vận chuyển\n\nMiễn phí từ 500.000đ."),
        "txt",
        "a" * 64,
    )
    parsed_b = ParsedDocument(
        MarkdownDocument("policy.txt", "Vận chuyển", "# Vận chuyển\n\nMiễn phí từ 1.000.000đ."),
        "txt",
        "b" * 64,
    )

    document_a = documents.add_document(dataset_a.dataset_id, parsed_a, FixtureEmbedder())
    document_b = documents.add_document(dataset_b.dataset_id, parsed_b, FixtureEmbedder())
    assert default_chunk > 0
    chunks_a = documents.retrieve([1.0, 0.0, *([0.0] * 766)], dataset_id=dataset_a.dataset_id)
    chunks_b = documents.retrieve([1.0, 0.0, *([0.0] * 766)], dataset_id=dataset_b.dataset_id)

    assert {chunk.content for chunk in chunks_a} == {"Miễn phí từ 500.000đ."}
    assert {chunk.content for chunk in chunks_b} == {"Miễn phí từ 1.000.000đ."}
    assert set(chunks_a).isdisjoint(chunks_b)
    assert {chunk.content for chunk in documents.retrieve([1.0, 0.0, *([0.0] * 766)])} == {
        "Default corpus must stay isolated."
    }
    with pytest.raises(ResourceConflictError):
        documents.add_document(dataset_a.dataset_id, parsed_a, FixtureEmbedder())

    with Session(management_database) as session:
        before_documents = session.scalar(select(func.count()).select_from(DatasetDocument))
        before_chunks = session.scalar(select(func.count()).select_from(DocumentChunk))
    failed = ParsedDocument(
        MarkdownDocument("failure.txt", "Failure", "# Failure\n\nRollback me."),
        "txt",
        "c" * 64,
    )
    with pytest.raises(RuntimeError, match="embedding unavailable"):
        documents.add_document(dataset_a.dataset_id, failed, FailingEmbedder())
    with Session(management_database) as session:
        assert session.scalar(select(func.count()).select_from(DatasetDocument)) == before_documents
        assert session.scalar(select(func.count()).select_from(DocumentChunk)) == before_chunks

    documents.delete_document(dataset_a.dataset_id, document_a.document_id)
    with pytest.raises(ResourceNotFoundError):
        documents.get_document(dataset_a.dataset_id, document_a.document_id)
    assert documents.get_document(dataset_b.dataset_id, document_b.document_id)
