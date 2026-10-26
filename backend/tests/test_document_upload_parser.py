"""Safety and format coverage for dataset document uploads."""

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, StreamObject

from app.core.errors import UploadValidationError
from app.rag.upload_parser import parse_document_file


def _docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_heading("Chính sách", level=1)
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_bytes(text: str | None = None, *, encrypted: bool = False) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text is not None:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_ref = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
        )
        stream = StreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("secret")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


@pytest.mark.parametrize("extension", ["txt", "md"])
def test_parses_utf8_text_and_sanitizes_filename(extension: str) -> None:
    parsed = parse_document_file(
        f"../../chính-sách.{extension}",
        "# Đổi trả\n\nKhách hàng được đổi trong 14 ngày.".encode(),
    )

    assert parsed.document.source_file == f"chính-sách.{extension}"
    assert parsed.document.title == "Đổi trả"
    assert "14 ngày" in parsed.document.content


def test_parses_docx_without_extracting_archive_to_disk() -> None:
    parsed = parse_document_file("quy-dinh.docx", _docx_bytes("Miễn phí vận chuyển từ 500.000đ."))

    assert parsed.file_type == "docx"
    assert parsed.document.title == "Chính sách"
    assert "500.000đ" in parsed.document.content


def test_parses_text_pdf_and_rejects_scanned_or_encrypted_pdf() -> None:
    parsed = parse_document_file("policy.pdf", _pdf_bytes("Return within 14 days"))

    assert "Return within 14 days" in parsed.document.content
    with pytest.raises(UploadValidationError, match="OCR"):
        parse_document_file("scan.pdf", _pdf_bytes())
    with pytest.raises(UploadValidationError, match="mã hóa"):
        parse_document_file("secret.pdf", _pdf_bytes(encrypted=True))


def test_pdf_replaces_extracted_nul_before_chunking_without_joining_words():
    import hashlib

    from app.rag.chunking import chunk_document

    content = _pdf_bytes("Service\x00handbook: fee 0.00 and 125000.50")
    parsed = parse_document_file("handbook.pdf", content)
    assert "\x00" not in parsed.document.content
    assert "Service handbook: fee 0.00 and 125000.50" in parsed.document.content
    assert parsed.content_hash == hashlib.sha256(content).hexdigest()
    chunks = chunk_document(parsed.document)
    assert chunks and all("\x00" not in chunk.content for chunk in chunks)


def test_pdf_preserves_unicode_newlines_tabs_and_numeric_text(monkeypatch):
    from types import SimpleNamespace

    monkeypatch.setattr(
        "app.rag.upload_parser.PdfReader",
        lambda *args, **kwargs: SimpleNamespace(
            is_encrypted=False,
            pages=[
                SimpleNamespace(
                    extract_text=lambda: "Phòng Gia đình\x00 Đủ chỗ\nPhí\t0.00 và 125000.50"
                )
            ],
        ),
    )
    parsed = parse_document_file("hướng-dẫn.pdf", b"%PDF-test")
    assert "Phòng Gia đình  Đủ chỗ\nPhí\t0.00 và 125000.50" in parsed.document.content


def test_pdf_with_only_nul_remains_unreadable():
    with pytest.raises(UploadValidationError, match="OCR"):
        parse_document_file("empty.pdf", _pdf_bytes("\x00\x00"))


@pytest.mark.parametrize("extension", ["md", "txt"])
def test_nul_in_source_text_remains_rejected(extension):
    with pytest.raises(UploadValidationError, match="null"):
        parse_document_file("invalid." + extension, b"text\x00text")


@pytest.mark.parametrize(
    ("filename", "content", "message"),
    [
        ("empty.txt", b"", "trống"),
        ("bad.txt", b"\xff\xfe", "UTF-8"),
        ("bad.pdf", b"not a pdf", "chữ ký PDF"),
        ("bad.docx", b"PK broken", "DOCX"),
        ("heading.md", "# Chỉ có tiêu đề".encode(), "chỉ có tiêu đề"),
        ("legacy.doc", b"legacy", "PDF, DOCX, TXT hoặc MD"),
    ],
)
def test_rejects_empty_corrupt_and_unsupported_documents(
    filename: str, content: bytes, message: str
) -> None:
    with pytest.raises(UploadValidationError, match=message):
        parse_document_file(filename, content)


def test_rejects_oversized_and_zip_bomb_metadata() -> None:
    with pytest.raises(UploadValidationError, match="giới hạn"):
        parse_document_file("large.txt", b"a" * 11, max_bytes=10)

    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("word/document.xml", b"x" * 100_000)
    with pytest.raises(UploadValidationError, match="tỷ lệ nén"):
        parse_document_file("bomb.docx", output.getvalue())


def test_docx_table_rules_are_extracted_in_document_order():
    import io

    from docx import Document

    document = Document()
    document.add_heading("Warehouse KPI", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Rule"
    table.cell(0, 1).text = "Threshold"
    table.cell(1, 0).text = "GV-KPI-KHO-08"
    table.cell(1, 1).text = "At least 8 delivered orders, not all statuses"
    document.add_paragraph("End of warehouse rule")
    stream = io.BytesIO()
    document.save(stream)
    parsed = parse_document_file("kpi.docx", stream.getvalue()).document.content
    assert "GV-KPI-KHO-08" in parsed
    assert "At least 8 delivered orders, not all statuses" in parsed
    assert parsed.index("Warehouse KPI") < parsed.index("Rule | Threshold")
    assert parsed.index("GV-KPI-KHO-08") < parsed.index("End of warehouse rule")


def test_docx_table_content_is_subject_to_extraction_limit(monkeypatch):
    import io

    from docx import Document

    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "x" * 101
    stream = io.BytesIO()
    document.save(stream)
    monkeypatch.setattr("app.rag.upload_parser.MAX_EXTRACTED_CHARACTERS", 100)
    with pytest.raises(UploadValidationError, match="giới hạn"):
        parse_document_file("bounded.docx", stream.getvalue())


def test_docx_nested_table_is_refused_instead_of_silently_losing_rules():
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Hidden rule"
    output = BytesIO()
    document.save(output)
    with pytest.raises(UploadValidationError, match="bảng lồng"):
        parse_document_file("nested.docx", output.getvalue())
