"""Generate synthetic maintenance source files without contacting application services."""

from __future__ import annotations

import argparse
import csv
import random
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from maintenance_documents import DEVICE_TYPES, HANDBOOK, PROCEDURE, SERVICE_GROUPS

SEED = 2601026
AS_OF = date(2026, 10, 26)
POLICY = (
    "# KPI trung tâm\n\n"
    "Mỗi trung tâm đạt KPI nếu tổng phi_bao_tri_vnd của các phiếu bảo trì thuộc trung tâm đó "
    "ít nhất 8000000 VND. Tính tất cả trạng thái phiếu bảo trì, không chỉ trạng thái Hoàn thành. "
    "Phân biệt từng trung tâm theo mã ma_trung_tam. "
    "SUM bỏ qua giá trị NULL, giữ số 0 là giá trị hợp lệ.\n"
)


def money(cents: int) -> str:
    return f"{cents // 100}.{cents % 100:02d}"


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    with path.open("x", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def create_tables(output: Path) -> None:
    rng = random.Random(SEED)
    places = [
        ("Hà Nội", "Hoàn Kiếm", "Bắc"),
        ("Hà Nội", "Cầu Giấy", "Bắc"),
        ("Hải Phòng", "Lê Chân", "Bắc"),
        ("Hải Phòng", "Hồng Bàng", "Bắc"),
        ("Quảng Ninh", "Hạ Long", "Bắc"),
        ("Bắc Ninh", "Kinh Bắc", "Bắc"),
        ("Thái Nguyên", "Sông Công", "Bắc"),
        ("Lào Cai", "Yên Bái", "Bắc"),
        ("Phú Thọ", "Việt Trì", "Bắc"),
        ("Ninh Bình", "Hoa Lư", "Bắc"),
        ("Thanh Hóa", "Lam Sơn", "Trung"),
        ("Nghệ An", "Vinh", "Trung"),
        ("Hà Tĩnh", "Thành Sen", "Trung"),
        ("Quảng Trị", "Đông Hà", "Trung"),
        ("Huế", "Phú Xuân", "Trung"),
        ("Đà Nẵng", "Hải Châu", "Trung"),
        ("Đà Nẵng", "Sơn Trà", "Trung"),
        ("Quảng Ngãi", "Cẩm Thành", "Trung"),
        ("Gia Lai", "Quy Nhơn", "Trung"),
        ("Khánh Hòa", "Nha Trang", "Trung"),
        ("TP Hồ Chí Minh", "Thủ Đức", "Nam"),
        ("TP Hồ Chí Minh", "Tân Bình", "Nam"),
        ("TP Hồ Chí Minh", "Phú Mỹ", "Nam"),
        ("Đồng Nai", "Biên Hòa", "Nam"),
        ("Tây Ninh", "Tân An", "Nam"),
        ("Cần Thơ", "Ninh Kiều", "Nam"),
        ("Cần Thơ", "Cái Răng", "Nam"),
        ("An Giang", "Long Xuyên", "Nam"),
        ("Vĩnh Long", "Long Châu", "Nam"),
        ("Cà Mau", "Tân Thành", "Nam"),
    ]
    centers = [
        {
            "ma_trung_tam": f"TT{i:02d}",
            "ten_trung_tam": f"Trung tâm {name}",
            "thanh_pho": city,
            "khu_vuc": region,
            "ngay_hoat_dong": (date(2023, 1, 1) + timedelta(days=i * 19)).isoformat(),
            "loai_trung_tam": ["Tổng hợp", "Chuyên dụng", "Lưu động"][i % 3],
            "trang_thai_hoat_dong": "Tạm ngưng" if i in (8, 23) else "Đang hoạt động",
        }
        for i, (city, name, region) in enumerate(places, 1)
    ]
    services = []
    for i in range(60):
        group, device = SERVICE_GROUPS[i % 6], DEVICE_TYPES[(i // 6) % 5]
        level = ["Cơ bản", "Nâng cao", "Chuyên sâu"][(i // 5) % 3]
        frequency = "định kỳ" if i < 30 else "theo yêu cầu"
        services.append(
            {
                "ma_dich_vu": f"DV{i + 1:02d}",
                "ten_dich_vu": f"{group} {device.lower()} {frequency}",
                "nhom_dich_vu": group,
                "loai_thiet_bi": device,
                "muc_do_ky_thuat": level,
                "thoi_gian_du_kien_gio": money(rng.randint(50, 600)),
                "phi_tham_khao_vnd": money(10000000 + i * 475031 + (50 if i % 3 else 0)),
                "trang_thai_dich_vu": "Tạm dừng" if i in (21, 47) else "Đang cung cấp",
            }
        )
    totals = [450000000 + i * 19000000 + i * 37 for i in range(13)]
    totals += [799500025, 799999999, 800000000, 800000001]
    totals += [820000000 + i * 31000000 + i * 113 for i in range(13)]
    assert len(totals) == 30
    rng.shuffle(totals)
    facts = []
    for center_index, target in enumerate(totals, 1):
        statuses = ["Hoàn thành"] * 60 + ["Đang xử lý"] * 20 + ["Mới tiếp nhận"] * 10 + ["Hủy"] * 10
        rng.shuffle(statuses)
        zero = {i for i, status in enumerate(statuses) if status == "Hủy"}
        charged_cancellations = rng.sample(sorted(zero), 3)
        zero.difference_update(charged_cancellations)
        zero.update(rng.sample([i for i, s in enumerate(statuses) if s == "Hoàn thành"], 2))
        weights = {i: rng.randint(500, 3500) for i in range(100) if i not in zero}
        denominator = sum(weights.values())
        fees = {i: target * weight // denominator for i, weight in weights.items()}
        residue = target - sum(fees.values())
        for i in list(fees)[:residue]:
            fees[i] += 1
        assert sum(fees.values()) == target
        for i, status in enumerate(statuses):
            sequence = (center_index - 1) * 100 + i + 1
            received = date(2026, 1, 1) + timedelta(days=rng.randrange(292))
            completed = (
                received + timedelta(days=rng.randrange(7)) if status == "Hoàn thành" else None
            )
            service_index = sequence - 1 if sequence <= 60 else rng.randrange(60)
            # Bias creates meaningful popularity differences while retaining every service.
            if sequence > 60 and sequence % 4 == 0:
                service_index = rng.randrange(10)
            facts.append(
                {
                    "ma_phieu": f"PT{sequence:06d}",
                    "ma_trung_tam": f"TT{center_index:02d}",
                    "ma_dich_vu": f"DV{service_index + 1:02d}",
                    "ngay_tiep_nhan": received.isoformat(),
                    "ngay_hoan_thanh": completed.isoformat() if completed else "",
                    "trang_thai": status,
                    "phi_bao_tri_vnd": money(fees.get(i, 0)),
                    "so_gio_xu_ly": ""
                    if status == "Mới tiếp nhận"
                    else "0.00"
                    if status == "Hủy"
                    else money(rng.randint(25, 850)),
                    "muc_do_uu_tien": ["Thấp", "Bình thường", "Cao", "Khẩn"][rng.randrange(4)],
                    "kenh_tiep_nhan": ["Cổng dịch vụ", "Điện thoại", "Trực tiếp", "Đối tác"][
                        rng.randrange(4)
                    ],
                    "ghi_chu": ""
                    if sequence % 7 == 0
                    else [
                        "Đã đối chiếu mã thiết bị; hồ sơ giả lập",
                        "Chờ xác nhận phạm vi công việc",
                        "Kiểm tra tại chỗ, ghi nhận kết quả đo",
                        "Đã xác nhận phạm vi dịch vụ",
                        "Đối soát phí thực tế theo phiếu",
                    ][sequence % 5],
                }
            )
    rng.shuffle(facts)
    for name, rows in zip(
        ("01_trung_tam.csv", "02_dich_vu.csv", "03_phieu_bao_tri.csv"),
        (centers, services, facts),
        strict=True,
    ):
        write_csv(output / name, rows)


def create_docx(output: Path) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(8), 1.12
    for name, size in (("Title", 23), ("Heading 1", 17), ("Heading 2", 12)):
        style = document.styles[name]
        style.font.name, style.font.size, style.font.color.rgb = (
            "Arial",
            Pt(size),
            RGBColor(0, 0, 0),
        )
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
    properties = document.core_properties
    properties.author = "ShopQuery AI"
    properties.title = "Quy trình bảo trì thiết bị"
    properties.created = properties.modified = datetime(2026, 10, 26, 0, 0)
    for index, (title, sections) in enumerate(PROCEDURE):
        if index:
            document.add_page_break()
        else:
            document.add_paragraph("Quy trình bảo trì thiết bị", "Title")
        document.add_paragraph(title, "Heading 1")
        for subtitle, body in sections:
            document.add_paragraph(subtitle, "Heading 2")
            document.add_paragraph(body)
    path = output / "05_quy_trinh_bao_tri.docx"
    buffer = BytesIO()
    document.save(buffer)
    with ZipFile(buffer) as archive:
        members = [(entry.filename, archive.read(entry)) for entry in archive.infolist()]
    with path.open("xb") as stream, ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        for name, data in members:
            info = ZipInfo(name, (2026, 10, 26, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            archive.writestr(info, data)


def create_pdf(output: Path, regular_font: Path, bold_font: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    for name, file in (("Body", regular_font), ("Heading", bold_font)):
        pdfmetrics.registerFont(TTFont(name, str(file)))
    style = ParagraphStyle(
        "Body", fontName="Body", fontSize=11, leading=16, spaceAfter=10, alignment=TA_LEFT
    )
    heading = ParagraphStyle(
        "Heading", parent=style, fontName="Heading", fontSize=18, leading=23, spaceAfter=16
    )
    subheading = ParagraphStyle(
        "Subheading",
        parent=style,
        fontName="Heading",
        fontSize=12,
        leading=17,
        spaceBefore=5,
        spaceAfter=5,
    )
    title_style = ParagraphStyle("Title", parent=heading, fontSize=23, leading=29)
    story = []
    for index, (title, sections) in enumerate(HANDBOOK):
        if index:
            story.append(PageBreak())
        else:
            story.extend([Paragraph("Sổ tay dịch vụ bảo trì", title_style), Spacer(1, 8)])
        story.append(Paragraph(title, heading))
        for subtitle, body in sections:
            story.extend([Paragraph(subtitle, subheading), Paragraph(body, style)])

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Body", 9)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(45, 27, "Sổ tay dịch vụ bảo trì thiết bị")
        canvas.drawRightString(A4[0] - 45, 27, str(doc.page))
        canvas.restoreState()

    # Exclusive creation protects existing files even when this helper is called directly.
    with (output / "06_so_tay_dich_vu.pdf").open("xb") as stream:
        document = SimpleDocTemplate(
            stream,
            pagesize=A4,
            leftMargin=45,
            rightMargin=45,
            topMargin=39,
            bottomMargin=45,
            title="Sổ tay dịch vụ bảo trì",
            author="ShopQuery AI",
            invariant=1,
        )
        document.build(story, onFirstPage=footer, onLaterPages=footer)


def generate(output: Path, regular_font: Path, bold_font: Path) -> None:
    """Create six synthetic files in a new, explicitly named directory; no services."""
    if not output.is_absolute():
        raise ValueError("Output must be an explicit absolute directory")
    if output.exists() or output.is_symlink():
        raise FileExistsError("Output directory already exists; nothing was overwritten")
    output = output.resolve()
    if not output.parent.is_dir():
        raise ValueError("Output parent must already exist")
    for font in (regular_font, bold_font):
        if not font.is_absolute() or not font.is_file():
            raise ValueError("Both font arguments must point to existing absolute TTF files")
    # Import and validate dependencies/fonts before creating any output directory.
    from docx import Document  # noqa: F401
    from reportlab.pdfbase.ttfonts import TTFont

    TTFont("PreflightBody", str(regular_font))
    TTFont("PreflightHeading", str(bold_font))
    output.mkdir()  # No exist_ok: a concurrent creator must cause a safe failure.
    create_tables(output)
    with (output / "04_KPI_trung_tam.md").open("x", encoding="utf-8") as stream:
        stream.write(POLICY)
    create_docx(output)
    create_pdf(output, regular_font, bold_font)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--font", type=Path, required=True, help="Absolute Unicode regular TTF")
    parser.add_argument("--bold-font", type=Path, required=True, help="Absolute Unicode bold TTF")
    args = parser.parse_args()
    try:
        generate(args.output, args.font, args.bold_font)
    except (ValueError, FileExistsError, ImportError, OSError) as exc:
        parser.exit(1, f"Cannot generate synthetic bundle: {exc}\n")
    print(f"Created six synthetic source files in {args.output}")


if __name__ == "__main__":
    main()
